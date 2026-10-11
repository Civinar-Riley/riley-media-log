# -*- coding: utf-8 -*-
"""
把 Excel 存量记录转换为 data/import.sql，供 wrangler d1 execute 导入。

规则（与设计共识一致）：
- 只导入 5 个分类表（动漫/电影/剧集/小说/综艺），「未知待定」不导入
- 小说表内「类型=漫画」的行归类为漫画(manga)，其余为小说(novel)
- 进度「已阅」或空 → 看过(done)；数字/其他文本 → 在看(watching) + 进度文本
- 地区按空格和 / 拆分为数组；类型标签按 / 拆分
- 小说表的「状态」→ 完结状态；「作者」→ 主创
- 标题+年份+类别 查重：疑似重复照常导入，备注前加 [疑似重复] 标记
- 拼音与首字母由 pypinyin 计算（用于搜索与排序）
"""
import json
import re
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

import openpyxl
from pypinyin import lazy_pinyin, Style

XLSX = Path(__file__).resolve().parent.parent / "电影动漫小说剧集.xlsm"
OUT_SQL = Path(__file__).resolve().parent.parent / "data" / "import.sql"
OUT_REPORT = Path(__file__).resolve().parent.parent / "data" / "import_report.txt"

SHEET_CATEGORY = {"动漫": "anime", "电影": "movie", "剧集": "series", "综艺": "variety", "小说": "novel"}

FINISH_MAP = {"完结": "finished", "连载": "ongoing", "烂尾": "ruined"}

SEASON_RE = re.compile(r"第\s*[（(]?\s*(\d{1,2})\s*[)）]?\s*季")
YEAR_RE = re.compile(r"(19|20)\d{2}")


def sql_str(v: str) -> str:
    return "'" + v.replace("'", "''") + "'"


def sql_json(arr: list) -> str:
    return sql_str(json.dumps(arr, ensure_ascii=False))


def pinyin_of(title: str) -> tuple[str, str]:
    full = "".join(lazy_pinyin(title, style=Style.NORMAL))
    full = re.sub(r"[^a-z0-9]", "", full.lower())
    initials = lazy_pinyin(title, style=Style.FIRST_LETTER)
    first = next((ch.upper() for ch in "".join(initials) if ch.isalpha() and ch.isascii()), "#")
    return full, first


def split_regions(v: str) -> list:
    if not v:
        return []
    parts = re.split(r"[/\s]+", v.strip())
    return [p.strip() for p in parts if p.strip()]


def split_tags(v: str) -> list:
    if not v:
        return []
    return [p.strip() for p in v.split("/") if p.strip()]


def parse_year(v) -> int | None:
    if v is None:
        return None
    m = YEAR_RE.search(str(v))
    return int(m.group(0)) if m else None


def parse_int(v) -> int | None:
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    m = re.search(r"\d+", s)
    return int(m.group(0)) if m else None


def map_status_progress(raw, category: str) -> tuple[str, str]:
    s = ("" if raw is None else str(raw)).strip()
    if s in ("", "已阅", "已读", "已看", "完"):
        return "done", ""
    n = re.search(r"\d+", s)
    if category in ("anime", "series") and n and re.fullmatch(r"\d+", s):
        return "watching", f"看到第 {int(s)} 集"
    return "watching", s


def detect_columns(header: list) -> dict:
    """根据表头文字猜列含义，返回 字段名 -> 列号"""
    col = {}
    for i, h in enumerate(header):
        if not h:
            continue
        h = str(h).strip()
        if i == 0 or any(k in h for k in ("电影", "剧集", "动漫", "综艺", "小说")):
            col.setdefault("title", i)
        if "时间" in h:
            col["year"] = i
        if "地区" in h:
            col["regions"] = i
        if "类型" in h:
            col["type"] = i
        if "作者" in h:
            col["creators"] = i
        if "集数" in h:
            col["episodes"] = i
        if h == "状态":
            col["finish"] = i
        if "进度" in h:
            col["progress"] = i
        if "备注" in h:
            col["notes"] = i
        if "观后感" in h or "读后感" in h:
            col["review"] = i
    return col


def cell(row, idx):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    now = datetime.now(timezone(timedelta(hours=8))).isoformat(timespec="seconds")
    rows_out = []
    stats = {}
    dupes = []
    seen_keys = set()

    for sheet_name, default_cat in SHEET_CATEGORY.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(values_only=True))
        # 跳过开头的空行
        while all_rows and not any(c not in (None, "") for c in all_rows[0]):
            all_rows.pop(0)
        if not all_rows:
            continue
        header = list(all_rows[0])
        cols = detect_columns(header)
        count = 0
        for raw in all_rows[1:]:
            if not any(c not in (None, "") for c in raw):
                continue
            title = str(cell(raw, cols.get("title")) or "").strip()
            if not title:
                continue

            category = default_cat
            if sheet_name == "小说":
                t = str(cell(raw, cols.get("type")) or "").strip()
                category = "manga" if t == "漫画" else "novel"

            year = parse_year(cell(raw, cols.get("year")))
            regions = split_regions(str(cell(raw, cols.get("regions")) or ""))
            tags = split_tags(str(cell(raw, cols.get("type") if sheet_name != "小说" else None) or ""))
            # 非小说表的「类型」列就是类型标签；小说表的类型列已用作分类
            if sheet_name == "小说":
                tags = []
            creators = split_tags(str(cell(raw, cols.get("creators")) or ""))
            episodes = parse_int(cell(raw, cols.get("episodes")))
            finish = FINISH_MAP.get(str(cell(raw, cols.get("finish")) or "").strip())
            progress_raw = cell(raw, cols.get("progress"))
            status, progress = map_status_progress(progress_raw, category)
            notes = str(cell(raw, cols.get("notes")) or "").strip()
            review = str(cell(raw, cols.get("review")) or "").strip()

            key = (title, year, category)
            if key in seen_keys:
                dupes.append(f"{sheet_name}:{title}({year})")
                notes = ("[疑似重复] " + notes).strip()
            seen_keys.add(key)

            season = None
            m = SEASON_RE.search(title)
            if m:
                season = int(m.group(1))

            full_py, first_py = pinyin_of(title)
            rows_out.append(
                f"INSERT INTO items (title, pinyin, pinyin_initial, category, year, season, "
                f"regions, tags, authors, status, progress, episode_count, rating, cover_url, "
                f"channel, finish_status, notes, review, tmdb_id, created_at, updated_at) VALUES ("
                f"{sql_str(title)}, {sql_str(full_py)}, {sql_str(first_py)}, {sql_str(category)}, "
                f"{year if year is not None else 'NULL'}, {season if season is not None else 'NULL'}, "
                f"{sql_json(regions)}, {sql_json(tags)}, {sql_json(creators)}, {sql_str(status)}, "
                f"{sql_str(progress)}, {episodes if episodes is not None else 'NULL'}, NULL, '', '', "
                f"{sql_str(finish) if finish else 'NULL'}, {sql_str(notes)}, {sql_str(review)}, NULL, "
                f"{sql_str(now)}, {sql_str(now)});"
            )
            count += 1
        stats[sheet_name] = count

    OUT_SQL.parent.mkdir(exist_ok=True)
    # 注意：不能用 BEGIN TRANSACTION 包裹——远程 D1（Durable Objects 存储）会拒绝该语句
    OUT_SQL.write_text("\n".join(rows_out) + "\n", encoding="utf-8")

    lines = [
        f"导入时间: {now}",
        f"总计: {len(rows_out)} 条",
        *[f"  {k}: {v} 条" for k, v in stats.items()],
        f"疑似重复: {len(dupes)} 条",
        *[f"  {d}" for d in dupes[:30]],
    ]
    OUT_REPORT.write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))
    print(f"\nSQL 已写入 {OUT_SQL}")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        import traceback

        traceback.print_exc()
        print(f"导入失败: {e}", file=sys.stderr)
        sys.exit(1)
